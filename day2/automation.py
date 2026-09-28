import openpyxl as xl
from openpyxl.chart import BarChart, Reference

def precess_workbook (filename):
    wb = xl.load_workbook(filename)
    sheet = wb['Transactions']


    for row in range(2, sheet.max_row + 1):
        cell = sheet.cell(row, 3)
        discount_price = cell.value * 0.9 # type: ignore
        discount_price_cell  = sheet.cell(row, 4)
        discount_price_cell.value = discount_price # type: ignore

    values = Reference(sheet, min_row=2, max_row=50, min_col=4, max_col=4)
    # values = Reference(sheet, min_row=2, max_row=sheet.max_row, min_col=4, max_col=4)

    chart = BarChart()
    chart.add_data(values)
    sheet.add_chart(chart, 'e2') # type: ignore


    wb.save(filename)